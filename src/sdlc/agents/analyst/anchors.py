"""Anchor reading for the analyst agent.

Anchors (DDS/TDS/...) are found on the Jira step tickets (ticket_anchors.py) and read IN
FULL — never through top-k retrieval. A Confluence page is read section by section (split
at h1-h3) with diagrams as image bytes, from the local corpus when synced, else live. An
attached file is read per format: xlsx -> one CSV section per sheet, docx -> sections split
at Heading 1-3 (plus tables and images), csv -> text, pdf -> passed through as a native
document for the model to read.

page_roles.yaml no longer drives anchor resolution; resolve_anchor_page_id below is still
used for the use case root page (where the plan is published).
"""
import csv
import io
import json
from dataclasses import dataclass, field
from pathlib import Path

import yaml
from bs4 import BeautifulSoup

from sdlc.agents.analyst.ticket_anchors import AttachmentSource, ConfluenceSource
from sdlc.ports import DocumentSource, ObjectStore, TicketSystem

_IMAGE_FORMATS = {"image/png": "png", "image/jpeg": "jpeg", "image/gif": "gif", "image/webp": "webp"}
_EXT_FORMATS = {"png": "png", "jpg": "jpeg", "jpeg": "jpeg", "gif": "gif", "webp": "webp"}


class AnchorNotConfirmed(Exception):
    def __init__(self, role: str, candidate_title: str | None):
        self.role = role
        self.candidates = [candidate_title] if candidate_title else []
        msg = f"anchor {role!r} has no page_id set in page_roles.yaml"
        if candidate_title:
            msg += f"; found a candidate by title: {candidate_title!r} — confirm by setting page_id"
        super().__init__(msg)


@dataclass
class ImageRef:
    format: str
    data: bytes


@dataclass
class DocumentRef:
    format: str  # a Converse document format the model reads natively (today: "pdf")
    data: bytes


@dataclass
class AnchorResult:
    page_id: str
    title: str
    url: str
    sections: list[tuple[str, str]]  # [(heading, text), ...]
    images: list[ImageRef] = field(default_factory=list)
    documents: list[DocumentRef] = field(default_factory=list)


def _page_roles(use_case: str) -> dict:
    return yaml.safe_load((Path("usecases") / use_case / "page_roles.yaml").read_text())["roles"]


def _split_sections(html: str) -> list[tuple[str, str]]:
    """Splits at h1-h3, matching how a human would read the page section by section."""
    soup = BeautifulSoup(html, "html.parser")
    sections: list[tuple[str, str]] = []
    heading, parts = "(untitled)", []
    for el in soup.find_all(True):
        if el.name in ("h1", "h2", "h3"):
            if parts:
                sections.append((heading, "\n".join(parts)))
            heading, parts = el.get_text(strip=True), []
        elif el.name in ("p", "li", "td", "th") and el.find(["h1", "h2", "h3"]) is None:
            text = el.get_text(" ", strip=True)
            if text:
                parts.append(text)
    if parts:
        sections.append((heading, "\n".join(parts)))
    if sections:
        return sections
    whole_text = soup.get_text(" ", strip=True)
    return [("(untitled)", whole_text)] if whole_text else []


def resolve_anchor_page_id(role: str, *, use_case: str, doc_source: DocumentSource) -> str:
    entry = _page_roles(use_case).get(role)
    if not entry:
        raise AnchorNotConfirmed(role, None)
    if entry.get("page_id"):
        return str(entry["page_id"])
    title_prefix = (entry.get("title_pattern") or "").rstrip("*")
    candidate = doc_source.find_page_by_title(title_prefix) if title_prefix else None
    raise AnchorNotConfirmed(role, candidate.title if candidate else None)


def _read_from_corpus(page_id: str, *, store: ObjectStore, data_dir: Path) -> AnchorResult | None:
    manifest = store.get_json("manifest.json") or {}
    entry = manifest.get(page_id)
    if not entry:
        return None
    page_dir = data_dir / "corpus" / entry["path"]
    html_path, meta_path = page_dir / "page.html", page_dir / "meta.json"
    if not html_path.exists() or not meta_path.exists():
        return None

    meta = json.loads(meta_path.read_text())
    images = []
    for filename in meta.get("attachments", []):
        fmt = _EXT_FORMATS.get(Path(filename).suffix.lower().lstrip("."))
        if fmt:
            images.append(ImageRef(format=fmt, data=(page_dir / "_attachments" / filename).read_bytes()))

    return AnchorResult(page_id=page_id, title=meta["title"], url=meta["url"],
                        sections=_split_sections(html_path.read_text()), images=images)


def _read_confluence_page(page_id: str, *, doc_source: DocumentSource, store: ObjectStore,
                           data_dir: Path) -> AnchorResult:
    from_corpus = _read_from_corpus(page_id, store=store, data_dir=data_dir)
    if from_corpus:
        return from_corpus

    page = doc_source.get_page(page_id)  # not yet synced locally — read live instead
    images = [
        ImageRef(format=_IMAGE_FORMATS[a.media_type], data=doc_source.download_attachment(a))
        for a in doc_source.get_attachments(page_id) if a.media_type in _IMAGE_FORMATS
    ]
    return AnchorResult(page_id=page.page_id, title=page.title, url=page.url,
                        sections=_split_sections(page.html), images=images)


def _read_xlsx(data: bytes) -> list[tuple[str, str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sections = []
    for sheet in workbook.worksheets:
        out = io.StringIO()
        writer = csv.writer(out, lineterminator="\n")
        for row in sheet.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                writer.writerow(["" if cell is None else cell for cell in row])
        sections.append((sheet.title, out.getvalue()))
    return sections


def _read_docx(data: bytes) -> tuple[list[tuple[str, str]], list[ImageRef]]:
    from docx import Document
    from docx.table import Table

    document = Document(io.BytesIO(data))
    sections: list[tuple[str, str]] = []
    heading, parts = "(untitled)", []
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            parts += [" | ".join(c.text.strip() for c in row.cells) for row in block.rows]
        elif block.style is not None and block.style.name in ("Heading 1", "Heading 2", "Heading 3"):
            if parts:
                sections.append((heading, "\n".join(parts)))
            heading, parts = block.text.strip(), []
        elif block.text.strip():
            parts.append(block.text.strip())
    if parts:
        sections.append((heading, "\n".join(parts)))

    images = [ImageRef(format=_IMAGE_FORMATS[rel.target_part.content_type], data=rel.target_part.blob)
              for rel in document.part.rels.values()
              if "/image" in rel.reltype and rel.target_part.content_type in _IMAGE_FORMATS]
    return sections, images


def _read_attachment(source: AttachmentSource, *, tickets: TicketSystem) -> AnchorResult:
    att = source.attachment
    data = tickets.download_attachment(att)
    result = AnchorResult(page_id=att.attachment_id, title=att.title, url=att.download_url, sections=[])
    if source.format == "xlsx":
        result.sections = _read_xlsx(data)
    elif source.format == "docx":
        result.sections, result.images = _read_docx(data)
    elif source.format == "csv":
        result.sections = [(att.title, data.decode("utf-8-sig"))]
    elif source.format == "pdf":
        result.documents = [DocumentRef(format="pdf", data=data)]
    else:
        raise ValueError(f"unsupported anchor attachment format {source.format!r}")
    return result


def read_anchor(source: ConfluenceSource | AttachmentSource, *, doc_source: DocumentSource,
                 tickets: TicketSystem, store: ObjectStore, data_dir: str | Path) -> AnchorResult:
    if isinstance(source, ConfluenceSource):
        return _read_confluence_page(source.page_id, doc_source=doc_source, store=store,
                                      data_dir=Path(data_dir))
    return _read_attachment(source, tickets=tickets)
