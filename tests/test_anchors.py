from pathlib import Path

import io

import pytest

from sdlc.adapters.fake import FakeDocumentSource, FakeObjectStore, FakeTicketSystem
from sdlc.agents.analyst.anchors import AnchorNotConfirmed, read_anchor, resolve_anchor_page_id
from sdlc.agents.analyst.ticket_anchors import AttachmentSource, ConfluenceSource
from sdlc.ports import Attachment

UC = "demo_order_fulfilment"  # uses the real usecases/demo_order_fulfilment/page_roles.yaml


def _read(source, *, doc=None, store=None, tickets=None, data_dir="/tmp"):
    return read_anchor(source, doc_source=doc or FakeDocumentSource(), tickets=tickets or FakeTicketSystem(),
                        store=store or FakeObjectStore(), data_dir=data_dir)


def test_missing_page_id_raises_with_no_candidate(monkeypatch):
    monkeypatch.setattr(
        "sdlc.agents.analyst.anchors._page_roles",
        lambda use_case: {"data_design_solution": {"title_pattern": "Data Design Solution*",
                                                     "page_id": None}},
    )
    doc = FakeDocumentSource()
    with pytest.raises(AnchorNotConfirmed) as exc:
        resolve_anchor_page_id("data_design_solution", use_case=UC, doc_source=doc)
    assert exc.value.role == "data_design_solution"
    assert exc.value.candidates == []


def test_missing_page_id_surfaces_a_title_candidate(monkeypatch):
    # title_pattern "Data Design Solution*"; the fake's find_page_by_title is an exact
    # match, so the title must match it exactly once the trailing '*' is stripped.
    monkeypatch.setattr(
        "sdlc.agents.analyst.anchors._page_roles",
        lambda use_case: {"data_design_solution": {"title_pattern": "Data Design Solution*",
                                                     "page_id": None}},
    )
    doc = FakeDocumentSource()
    doc.add_page("Data Design Solution", page_id="9")
    with pytest.raises(AnchorNotConfirmed) as exc:
        resolve_anchor_page_id("data_design_solution", use_case=UC, doc_source=doc)
    assert exc.value.candidates == ["Data Design Solution"]


def test_reads_from_local_corpus_when_synced(tmp_path: Path):
    store = FakeObjectStore()
    page_dir = tmp_path / "corpus" / "FAKE" / "data-design-solution__42"
    page_dir.mkdir(parents=True)
    (page_dir / "page.html").write_text(
        "<h2>Layers</h2><p>Raw, Curated, Presentation.</p>"
        "<h2>Refresh schedule</h2><p>Daily by 6am.</p>")
    (page_dir / "meta.json").write_text(
        '{"title": "Data Design Solution", "url": "fake://42", "attachments": ["diagram.png"]}')
    (page_dir / "_attachments").mkdir()
    (page_dir / "_attachments" / "diagram.png").write_bytes(b"PNGDATA")
    store.put_json("manifest.json", {"42": {"version": 1, "path": "FAKE/data-design-solution__42"}},
                    dry_run=False)

    result = _read(ConfluenceSource("42", "DDS"), store=store, data_dir=tmp_path)

    assert result.title == "Data Design Solution"
    assert result.sections == [
        ("Layers", "Raw, Curated, Presentation."),
        ("Refresh schedule", "Daily by 6am."),
    ]
    assert len(result.images) == 1
    assert result.images[0].format == "png"
    assert result.images[0].data == b"PNGDATA"


def test_falls_back_to_live_read_when_not_in_corpus():
    doc = FakeDocumentSource()
    doc.add_page("Data Design Solution", page_id="42", html="<h1>Intro</h1><p>Hello.</p>")

    result = _read(ConfluenceSource("42", "DDS"), doc=doc)  # empty manifest -> not in corpus

    assert result.sections == [("Intro", "Hello.")]


def _attachment_read(name: str, fmt: str, data: bytes):
    tickets = FakeTicketSystem()
    att = Attachment("a1", name, "x/y", "https://dl/a1")
    tickets.attachment_bytes["a1"] = data
    return _read(AttachmentSource(att, fmt), tickets=tickets)


def test_reads_xlsx_one_csv_section_per_sheet():
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Mappings"
    wb.active.append(["src", "tgt"])
    wb.active.append(["a", 1])
    wb.active.append([None, None])  # blank row is skipped
    wb.create_sheet("Entities").append(["order", None, "id"])
    buf = io.BytesIO()
    wb.save(buf)

    result = _attachment_read("DDS.xlsx", "xlsx", buf.getvalue())

    assert result.sections == [("Mappings", "src,tgt\na,1\n"), ("Entities", "order,,id\n")]


def test_reads_docx_sections_tables_and_images():
    from docx import Document
    from PIL import Image

    png = io.BytesIO()
    Image.new("RGB", (2, 2)).save(png, format="PNG")
    png.seek(0)
    d = Document()
    d.add_paragraph("preamble")
    d.add_heading("Layers", level=1)
    d.add_paragraph("Raw, Curated.")
    d.add_heading("Schedule", level=2)
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "daily", "6am"
    d.add_picture(png)
    buf = io.BytesIO()
    d.save(buf)

    result = _attachment_read("TDS.docx", "docx", buf.getvalue())

    assert result.sections == [("(untitled)", "preamble"), ("Layers", "Raw, Curated."),
                                ("Schedule", "daily | 6am")]
    assert [i.format for i in result.images] == ["png"]


def test_reads_csv_as_one_section():
    result = _attachment_read("map.csv", "csv", "\ufeffa,b\n1,2\n".encode("utf-8"))
    assert result.sections == [("map.csv", "a,b\n1,2\n")]


def test_pdf_is_passed_through_as_a_native_document():
    result = _attachment_read("TDS.pdf", "pdf", b"%PDF-1.7")
    assert result.sections == [] and [(d.format, d.data) for d in result.documents] == [("pdf", b"%PDF-1.7")]


def test_split_sections_does_not_repeat_text_wrapped_in_nested_blocks():
    from sdlc.agents.analyst.anchors import _split_sections

    html = ("<h2>Layers</h2><table><tr><td><p>cell</p></td><th><p>head</p></th></tr></table>"
            "<ul><li><p>one</p></li><li>two <ul><li><p>sub</p></li></ul></li></ul><!-- hidden -->")
    assert _split_sections(html) == [("Layers", "cell\nhead\none\ntwo\nsub")]


def test_split_sections_keeps_text_that_sits_beside_a_nested_block():
    from sdlc.agents.analyst.anchors import _split_sections

    html = "<h2>H</h2><table><tr><td>before <p>inside</p> after</td></tr></table>"
    assert _split_sections(html) == [("H", "before after\ninside")]
